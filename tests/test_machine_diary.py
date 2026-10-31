"""Machine-only exports and journal tamper rejection; isolated mock database."""
import io
import os
import sys
from copy import deepcopy
from pathlib import Path

os.environ.setdefault('MONGO_URL', 'mongodb://localhost:27017')
os.environ.setdefault('DB_NAME', 'machine_diary_test')
os.environ.setdefault('JWT_SECRET', 'machine-diary-test-only')
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))

import pytest
from fastapi import HTTPException
from mongomock_motor import AsyncMongoMockClient
from openpyxl import load_workbook
import fitz
import routes_data
import routes_records
import server
from exports import diary_pdf
from machine_diary import diary_projection, verified_diary_records, verified_diary_attachment
from record_integrity import chain_hash, CHAIN_GENESIS


def record(**extra):
    return {'id': 'machine-entry', 'date': '2026-10-30', 'revision': 1, 'status': 'confirmed',
            'owner_id': 'operator', 'operator': 'Test operator', 'machine': 'TEST/1', 'site': 'Test yard',
            'work': 'Obsluha linky', 'billing': [{'start': '06:00', 'end': '17:00'}],
            'mthStart': '500', 'mthEnd': '502', 'service': [], 'faults': [],
            'attachments': [], 'author_name': 'Test operator', 'updatedAt': '2026-10-30T17:00:00Z', **extra}


async def seed(db, value):
    payload = deepcopy(value)
    digest = chain_hash(CHAIN_GENESIS, payload)
    current = {**payload, 'integrity_hash': digest, 'prev_hash': CHAIN_GENESIS}
    await db.records.insert_one(deepcopy(current))
    await db.record_versions.insert_one({'record_id': value['id'], 'revision': value['revision'],
        'payload': payload, 'chain': 1, 'prev_hash': CHAIN_GENESIS, 'hash': digest})
    return current


@pytest.mark.asyncio
async def test_projection_tamper_keeps_hash_but_blocks_export(monkeypatch):
    db = AsyncMongoMockClient().diary
    current = await seed(db, record())
    monkeypatch.setattr(routes_records, 'db', db)
    assert (await routes_records.verify_chain(current['id']))['ok']
    assert (await verified_diary_records(db, [current]))[0]['work'] == 'Obsluha linky'
    current['work'] = 'Silent overwrite'
    await db.records.update_one({'id': current['id']}, {'$set': {'work': current['work']}})
    assert not (await routes_records.verify_chain(current['id']))['ok']
    with pytest.raises(HTTPException) as error:
        await verified_diary_records(db, [current])
    assert error.value.status_code == 409


@pytest.mark.asyncio
async def test_old_revision_tamper_and_missing_history_are_rejected():
    db = AsyncMongoMockClient().diary
    current = await seed(db, record())
    await db.record_versions.update_one({'record_id': current['id']}, {'$set': {'payload.work': 'Tampered'}})
    with pytest.raises(HTTPException):
        await verified_diary_records(db, [current])
    with pytest.raises(HTTPException):
        await verified_diary_records(db, [record(id='no-history')])


def test_pdf_and_xlsx_do_not_leak_commute_or_employee_shift():
    r = record(vehiclesTo='COMMUTE-SECRET', vehiclesFrom='COMMUTE-SECRET', accommodation='HOTEL-SECRET',
               notes='EMPLOYEE-NOTES-SECRET', drive='1:45', driveMinutes=105,
               travel=[{'start': '06:00', 'end': '07:00', 'from': 'TRAVEL-SECRET', 'to': 'Elsewhere'}],
               arrival='ARRIVAL-SECRET', departure='DEPARTURE-SECRET', work='=1+1',
               handover='Machine handed to next operator')
    clean = diary_projection(r)
    assert not {'billing', 'travel', 'drive', 'notes', 'vehiclesTo', 'arrival', 'departure'} & clean.keys()
    with fitz.open(stream=diary_pdf([r]), filetype='pdf') as pdf:
        text = '\n'.join(page.get_text() for page in pdf)
    assert 'Machine handed to next operator' in text
    assert 'Čas celkem' not in text
    for forbidden in ['COMMUTE-SECRET', 'HOTEL-SECRET', 'TRAVEL-SECRET', 'ARRIVAL-SECRET', 'EMPLOYEE-NOTES-SECRET']:
        assert forbidden not in text
    book = load_workbook(io.BytesIO(routes_data.machine_logs_xlsx({'TEST/1': [r]}, {}, {}, {})))
    assert book.security.lockStructure
    for sheet in book:
        assert sheet.protection.sheet
        assert all(cell.protection.locked for row in sheet for cell in row)
    sheet = book.worksheets[1]
    headers = [cell.value for cell in sheet[4]]
    assert 'MTH rozdíl' in headers
    assert sheet.cell(5, headers.index('MTH rozdíl') + 1).value == 2
    for header in ['SPZ na pracoviště', 'Pracovní doba (úseky)', 'Celkem h', 'Přejezdy h', 'K fakturaci h']:
        assert header not in headers
    values = [cell.value for ws in book for row in ws for cell in row]
    assert '=1+1' in values
    assert all(cell.data_type != 'f' for ws in book for row in ws for cell in row)
    assert not any(value in values for value in ['COMMUTE-SECRET', 'HOTEL-SECRET', 'EMPLOYEE-NOTES-SECRET', 'TRAVEL-SECRET'])


@pytest.mark.asyncio
async def test_both_export_paths_default_to_confirmed_and_verify_snapshots(monkeypatch):
    db = AsyncMongoMockClient().diary
    actor = {'id': 'operator', 'name': 'Test operator', 'role': 'employee'}
    await db.users.insert_one(dict(actor))
    await seed(db, record())
    await seed(db, record(id='draft', status='draft'))
    monkeypatch.setattr(routes_data, 'db', db)
    monkeypatch.setattr(server, 'db', db)
    grouped, _, _ = await routes_data.machine_log_data(actor, {})
    assert [r['id'] for r in grouped['TEST/1']] == ['machine-entry']
    assert all('billing' not in r for r in grouped['TEST/1'])
    official = await server.export_diary(user=actor)
    with fitz.open(stream=official.body, filetype='pdf') as pdf:
        assert 'PRACOVNÍ VERZE' not in ''.join(p.get_text() for p in pdf)
    draft_export = await server.export_diary(user=actor, include_drafts=1)
    with fitz.open(stream=draft_export.body, filetype='pdf') as pdf:
        assert 'PRACOVNÍ VERZE' in ''.join(p.get_text() for p in pdf)
    await db.records.update_one({'id': 'machine-entry'}, {'$set': {'work': 'Projection overwrite'}})
    with pytest.raises(HTTPException):
        await routes_data.machine_log_data(actor, {})
    with pytest.raises(HTTPException):
        await server.export_diary(user=actor)


def test_changed_attachment_bytes_fail_instead_of_silently_exporting():
    import hashlib
    source = b'original photo bytes'
    attachment = {'sha256': hashlib.sha256(source).hexdigest()}
    assert verified_diary_attachment(attachment, {'storage_path': 'opaque'}, lambda _: (source, 'image/jpeg')) == source
    with pytest.raises(HTTPException) as error:
        verified_diary_attachment(attachment, {'storage_path': 'opaque'}, lambda _: (b'overwritten', 'image/jpeg'))
    assert error.value.status_code == 409
