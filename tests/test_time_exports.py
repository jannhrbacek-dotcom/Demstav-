"""Regression tests for time, personal journey and consistent exports; isolated DB."""
import asyncio
import io
import os
import sys
from types import SimpleNamespace
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "audit_only")
os.environ.setdefault("JWT_SECRET", "audit-test-only")
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

import pytest
from fastapi import HTTPException
from mongomock_motor import AsyncMongoMockClient
from openpyxl import load_workbook
import core
import routes_records
import routes_data
import routes_production
from production import clean_report
from worktime import day_union, time_breakdown


def record(**kw):
    return dict(id="audit-record", date="2026-10-28", machine="TEST-2",
                operator="Test operator", mthStart="14000", mthEnd="14002",
                billing=[dict(start="06:00", end="17:00")], **kw)


@pytest.fixture
def database(monkeypatch):
    db = AsyncMongoMockClient().audit_only
    for module in (core, routes_records, routes_data, routes_production):
        monkeypatch.setattr(module, "db", db)
    return db


def test_shift_categories_are_not_added_twice():
    bd = time_breakdown(record(service=[dict(start="11:00", end="13:00")],
                               travel=[dict(start="14:00", end="15:00")]))
    assert bd["totalHours"] == 10
    assert bd["serviceHours"] == 1.5
    assert bd["travelHours"] == 1
    assert bd["workHours"] == 7.5


def test_personal_drive_accepts_hours_minutes():
    errors, _ = routes_records.validate_record(record(drive="0:45"), True)
    assert not errors, errors


def test_overnight_multiple_machines_union():
    first = {**record(), "billing": [dict(start="22:00", end="02:00")]}
    second = {**record(), "billing": [dict(start="01:00", end="03:00")]}
    assert day_union([first, second])["totalHours"] == 5


@pytest.mark.asyncio
async def test_machine_xlsx_matches_approved_breaks(database):
    actor = dict(id="employee-a", name="Test A", role="employee")
    await database.users.insert_one(actor.copy())
    await database.records.insert_one(record(owner_id=actor["id"], status="confirmed"))
    await database.break_requests.insert_one(dict(employee_id=actor["id"], date="2026-10-28",
        status="approved", breaks=[], decision=dict(at="2026-10-28T18:00:00Z")))
    by_machine, users, machines = await routes_data.machine_log_data(actor, {})
    xlsx = routes_data.machine_logs_xlsx(by_machine, users, machines, {})
    sheet = load_workbook(io.BytesIO(xlsx)).worksheets[1]
    assert sheet.cell(5, 12).value == 11


@pytest.mark.asyncio
async def test_timesheet_data_keeps_personal_journey(database):
    actor = dict(id="employee-a", name="Test A", role="employee")
    await database.users.insert_one(actor.copy())
    await database.records.insert_one(record(owner_id=actor["id"], status="confirmed", drive="45",
        vehiclesTo="TEST-PLATE", vehiclesFrom="TEST-PLATE"))
    rows, _, _ = await routes_data.timesheet_rows(actor, {})
    assert "Jízda" in rows[0]
    assert "SPZ na pracoviště" in rows[0]


def test_implausible_mth_requires_review():
    errors, _ = routes_records.validate_record({**record(), "mthEnd": "15000"}, True)
    assert errors


def test_driver_and_passenger_pdf_keep_own_duration_and_zero():
    import fitz
    from exports import timesheet_pdf
    from journey import journey_text
    groups = []
    for role in ("driver", "passenger"):
        r = record(owner_id=role, status="confirmed", revision=1, drive="0:45",
            driveMinutes=45, journeyRole=role, vehiclesTo="TEST-1", vehiclesFrom="TEST-2")
        groups.append(dict(employee_name=role, records=[r]))
    pdf = timesheet_pdf(groups, lambda _: None)
    doc = fitz.open(stream=pdf, filetype="pdf")
    text = "\n".join(page.get_text() for page in doc)
    assert text.count("0:45") >= 2
    assert "TEST-1" in text and "TEST-2" in text
    assert "neověřeno" in text
    assert journey_text(dict(driveMinutes=0)) == "0:00"
    assert journey_text({}) == ""
    assert journey_text(dict(drive="45")) == "neověřeno"
