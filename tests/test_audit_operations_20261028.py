"""Isolated audit acceptance tests. Strict xfails are confirmed OPEN defects, not fixes.

Run from repository root with the cloud test dependencies and pytest-xdist installed.
No production credentials, network calls, startup migrations or real MongoDB are used.
"""
import asyncio
import io
import os
import sys
from types import SimpleNamespace
from pathlib import Path

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "audit_only")
os.environ.setdefault("JWT_SECRET", "audit-test-only")
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

import pytest
from fastapi import HTTPException
from mongomock_motor import AsyncMongoMockClient
from openpyxl import load_workbook
import core
import routes_records
import routes_data
import routes_production
from production import clean_report
from worktime import day_union, time_breakdown


def record(**kw):
    return dict(id="audit-record", date="2026-10-28", machine="TEST-2",
                operator="Test operator", mthStart="14000", mthEnd="14002",
                billing=[dict(start="06:00", end="17:00")], **kw)


@pytest.fixture
def database(monkeypatch):
    db = AsyncMongoMockClient().audit_only
    for module in (core, routes_records, routes_data, routes_production):
        monkeypatch.setattr(module, "db", db)
    return db


def test_shift_categories_are_not_added_twice():
    bd = time_breakdown(record(service=[dict(start="11:00", end="13:00")],
                               travel=[dict(start="14:00", end="15:00")]))
    assert bd["totalHours"] == 10
    assert bd["serviceHours"] == 1.5
    assert bd["travelHours"] == 1
    assert bd["workHours"] == 7.5


@pytest.mark.xfail(strict=True, reason="AUD-01: personal journey duration rejected")
def test_personal_drive_accepts_hours_minutes():
    errors, _ = routes_records.validate_record(record(drive="0:45"), True)
    assert not errors, errors


@pytest.mark.xfail(strict=True, reason="AUD-02: union has no common overnight anchor")
def test_overnight_multiple_machines_union():
    first = {**record(), "billing": [dict(start="22:00", end="02:00")]}
    second = {**record(), "billing": [dict(start="01:00", end="03:00")]}
    assert day_union([first, second])["totalHours"] == 5


@pytest.mark.xfail(strict=True, reason="AUD-03: employee can select external service")
def test_interval_service_requires_actor_permission():
    assert not routes_records.can_record_interval_service(
        dict(id="ordinary", name="Test employee", role="employee", position="strojnik"), "external")


@pytest.mark.xfail(strict=True, reason="AUD-04: duplicate response precedes ownership check")
@pytest.mark.asyncio
async def test_record_retry_does_not_disclose_foreign_record(database):
    owner = dict(id="employee-a", name="Test A", role="employee")
    other = dict(id="employee-b", name="Test B", role="employee")
    body = dict(opId="audit-operation", record=record())
    await routes_records.save_record("audit-record", body, owner)
    with pytest.raises(HTTPException) as exc:
        await routes_records.save_record("audit-record", body, other)
    assert exc.value.status_code in (403, 404)


@pytest.mark.xfail(strict=True, reason="AUD-05: record replace lacks revision compare-and-swap")
@pytest.mark.asyncio
async def test_concurrent_record_save_returns_conflict(database, monkeypatch):
    actor = dict(id="employee-a", name="Test A", role="employee")
    await database.record_versions.create_index([("record_id", 1), ("revision", 1)], unique=True)
    await routes_records.save_record("audit-record", dict(opId="initial", record=record()), actor)
    collection = database.records
    original = collection.replace_one
    monkeypatch.setattr(routes_records, "db", SimpleNamespace(records=collection,
        record_versions=database.record_versions, attachments=database.attachments))
    reached = 0
    barrier = asyncio.Event()

    async def simultaneous(*args, **kwargs):
        nonlocal reached
        reached += 1
        if reached == 2:
            barrier.set()
        await asyncio.wait_for(barrier.wait(), timeout=3)
        return await original(*args, **kwargs)

    monkeypatch.setattr(collection, "replace_one", simultaneous)
    result = await asyncio.gather(*[
        routes_records.save_record("audit-record", dict(opId=f"writer-{i}", baseRevision=1,
            record=record(work=f"writer-{i}")), actor) for i in range(2)
    ], return_exceptions=True)
    assert sum(isinstance(r, dict) for r in result) == 1
    assert any(isinstance(r, HTTPException) and r.status_code == 409 for r in result), result


@pytest.mark.xfail(strict=True, reason="AUD-06: total tonnes remains a second manual truth")
def test_tonnage_is_derived_from_tickets():
    report = clean_report(dict(date="2026-10-28", tonnes="99", tickets=[
        dict(attachment_id="ticket-0001", tonnes="10,5"),
        dict(attachment_id="ticket-0002", tonnes="20")]))
    assert report["tonnes"] == 30.5


@pytest.mark.xfail(strict=True, reason="AUD-07: machine XLSX omits approved break hydration")
@pytest.mark.asyncio
async def test_machine_xlsx_matches_approved_breaks(database):
    actor = dict(id="employee-a", name="Test A", role="employee")
    await database.users.insert_one(actor.copy())
    await database.records.insert_one(record(owner_id=actor["id"], status="confirmed"))
    await database.break_requests.insert_one(dict(employee_id=actor["id"], date="2026-10-28",
        status="approved", breaks=[], decision=dict(at="2026-10-28T18:00:00Z")))
    by_machine, users, machines = await routes_data.machine_log_data(actor, {})
    xlsx = routes_data.machine_logs_xlsx(by_machine, users, machines, {})
    sheet = load_workbook(io.BytesIO(xlsx)).worksheets[1]
    assert sheet.cell(5, 12).value == 11


@pytest.mark.xfail(strict=True, reason="AUD-08: timesheet data export omits journey and plates")
@pytest.mark.asyncio
async def test_timesheet_data_keeps_personal_journey(database):
    actor = dict(id="employee-a", name="Test A", role="employee")
    await database.users.insert_one(actor.copy())
    await database.records.insert_one(record(owner_id=actor["id"], status="confirmed", drive="45",
        vehiclesTo="TEST-PLATE", vehiclesFrom="TEST-PLATE"))
    rows, _, _ = await routes_data.timesheet_rows(actor, {})
    assert "Jízda" in rows[0]
    assert "SPZ na pracoviště" in rows[0]


@pytest.mark.xfail(strict=True, reason="AUD-09: responsible IDs are informational only")
@pytest.mark.asyncio
async def test_unassigned_employee_cannot_create_shared_report(database):
    actor = dict(id="unassigned", name="Test employee", role="employee")
    await database.production_settings.insert_one(dict(id="main", responsible_ids=["assigned"]))
    with pytest.raises(HTTPException) as exc:
        await routes_production.save("report-00001", dict(opId="operation-1", report=dict(date="2026-10-28")), actor)
    assert exc.value.status_code == 403


@pytest.mark.xfail(strict=True, reason="AUD-10: implausible MTH delta has no validation")
def test_implausible_mth_requires_review():
    errors, _ = routes_records.validate_record({**record(), "mthEnd": "15000"}, True)
    assert errors
