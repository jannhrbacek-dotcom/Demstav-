"""
Verify machine_log XLSX structure
"""
import io
import requests
from openpyxl import load_workbook

BASE_URL = "https://markdown-notes-12.preview.emergentagent.com/api"

# Login as manager
print("Logging in as manager...")
resp = requests.post(f"{BASE_URL}/auth/login", json={"email": "vedouci@demstav.cz", "password": "Demstav2026!"})
assert resp.status_code == 200
token = resp.json()["access_token"]
print("✓ Logged in")

# Get machine_log XLSX
print("\nFetching machine_log XLSX...")
headers = {"Authorization": f"Bearer {token}"}
resp = requests.get(f"{BASE_URL}/exports/data", headers=headers, params={
    "dataset": "machine_log",
    "format": "xlsx"
})

assert resp.status_code == 200, f"Expected 200, got {resp.status_code}"
print("✓ Got XLSX response")

# Load workbook
wb = load_workbook(io.BytesIO(resp.content))
print(f"\n✓ Loaded workbook with {len(wb.sheetnames)} sheets")
print(f"  Sheets: {wb.sheetnames}")

# Verify first sheet is 'Přehled'
assert wb.sheetnames[0] == "Přehled", f"First sheet should be 'Přehled', got '{wb.sheetnames[0]}'"
print("\n✅ PASS: First sheet is 'Přehled'")

# Check overview sheet
overview = wb["Přehled"]
print(f"\n✓ Overview sheet title: {overview['A1'].value}")
print(f"✓ Overview has {overview.max_row} rows, {overview.max_column} columns")

# Find a machine sheet (not Přehled)
machine_sheets = [name for name in wb.sheetnames if name != "Přehled"]
print(f"\n✓ Found {len(machine_sheets)} machine sheets")

if machine_sheets:
    # Check first machine sheet
    sheet_name = machine_sheets[0]
    ws = wb[sheet_name]
    print(f"\n✓ Checking machine sheet: '{sheet_name}'")
    
    # Find header row (row with "Datum" in first column)
    header_row = None
    for row_idx in range(1, 10):
        if ws.cell(row_idx, 1).value == "Datum":
            header_row = row_idx
            break
    
    assert header_row is not None, "Could not find header row with 'Datum'"
    print(f"✓ Found header row at row {header_row}")
    
    # Get all headers
    headers = []
    for col_idx in range(1, 60):
        val = ws.cell(header_row, col_idx).value
        if val:
            headers.append(val)
        elif col_idx > len(headers) + 5:
            break
    
    print(f"\n✓ Found {len(headers)} columns")
    print(f"  First 10: {headers[:10]}")
    print(f"  Last 10: {headers[-10:]}")
    
    # Verify expected columns
    expected_cols = [
        "Datum", "Směna", "Zaměstnanec", "Objednatel", "Pracoviště", "Zakázka",
        "MTH začátek", "MTH konec", "MTH celkem",
        "Denní kontrola", "Kontrola příslušenství", "Kontrola vybavení",
        "Provozuschopný", "Závady (shrnutí)", "STOP závady",
        "Servis provedl", "Externí firma", "Popis servisu", "Popis opravy",
        "Pravidelný servis (interval)",
        "Nafta l", "Hydr. olej l", "Mot. olej l", "Jiná kapalina",
        "Předal (jméno)", "Poznámky",
        "Stav zápisu", "Revize", "Důvod opravy", "Autor", "ID zápisu", "Otisk SHA-256"
    ]
    
    missing = []
    for exp in expected_cols:
        if exp not in headers:
            missing.append(exp)
    
    if missing:
        print(f"\n❌ FAIL: Missing columns: {missing}")
    else:
        print(f"\n✅ PASS: All expected columns present")
    
    # Check column count (should be ~46)
    if 44 <= len(headers) <= 48:
        print(f"✅ PASS: Column count {len(headers)} is in expected range (44-48)")
    else:
        print(f"❌ FAIL: Column count {len(headers)} not in expected range (44-48)")
    
    # Check for data rows
    data_rows = ws.max_row - header_row
    print(f"\n✓ Machine sheet has {data_rows} data rows")
else:
    print("\n⚠ WARNING: No machine sheets found (no records in database)")

print("\n" + "=" * 60)
print("XLSX STRUCTURE VERIFICATION COMPLETE")
print("=" * 60)
